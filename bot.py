import os
import io
import re
import sys
import telebot
from flask import Flask, request, jsonify
from docx import Document
from docx.shared import Pt, RGBColor
from weasyprint import HTML

# ==========================================
# 1. कॉन्फ़िगरेशन
# ==========================================
BOT_TOKEN = os.getenv("BOT_TOKEN")
# यदि Render में WEBHOOK_URL नहीं डाला है, तो सीधे अपनी सर्विस का URL लिखें
WEBHOOK_URL = os.getenv("WEBHOOK_URL", "https://mcq-converter-bot.onrender.com")

if not BOT_TOKEN:
    sys.exit("❌ Error: कृपया Render Environment Variables में 'BOT_TOKEN' सेट करें!")

bot = telebot.TeleBot(BOT_TOKEN, threaded=False)
app = Flask(__name__)


# ==========================================
# 2. DOCX फ़ाइल से टेक्स्ट पढ़ने वाला फ़ंक्शन
# ==========================================
def extract_text_from_docx(file_bytes):
    """Word (.docx) फाइल से पूरा टेक्स्ट निकालता है (पैराग्राफ + टेबल्स)"""
    doc = Document(io.BytesIO(file_bytes))
    full_text = []

    # सामान्य पैराग्राफ्स पढ़ें
    for p in doc.paragraphs:
        if p.text.strip():
            full_text.append(p.text.strip())

    # अगर डेटा टेबल्स के अंदर हो तो उसे भी पढ़ें
    for table in doc.tables:
        for row in table.rows:
            row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_data:
                full_text.append(" | ".join(row_data))

    return "\n".join(full_text)


# ==========================================
# 3. आपकी फ़ाइल के अनुसार MCQ पार्सर
# ==========================================
def parse_mcqs(text):
    """
    Question, Options, Answer, Solution, Key Points
    आदि फ़ील्ड्स को संरचित (Structured) रूप में अलग करता है।
    """
    # 'Question:' के आधार पर प्रश्नों को अलग करें
    raw_blocks = re.split(r'\n(?=Question:)', text.strip(), flags=re.IGNORECASE)
    parsed_questions = []

    for block in raw_blocks:
        if not block.strip():
            continue

        q_match = re.search(r'Question:\s*(.*?)(?=\([a-d]\)|Answer:|$)', block, re.DOTALL | re.IGNORECASE)
        ans_match = re.search(r'Answer:\s*([a-d])', block, re.IGNORECASE)
        sol_match = re.search(r'Solution:\s*(.*?)(?=Key Points:|Positive Marks:|$)', block, re.DOTALL | re.IGNORECASE)
        kp_match = re.search(r'Key Points:\s*(.*?)(?=Positive Marks:|Question:|$)', block, re.DOTALL | re.IGNORECASE)
        
        # (a), (b), (c), (d) विकल्प निकालना
        options = re.findall(r'\(([a-d])\)\s*([^(\n]+)', block, re.IGNORECASE)

        question_text = q_match.group(1).strip() if q_match else ""
        
        if question_text:
            parsed_questions.append({
                "question": question_text,
                "options": options,
                "answer": ans_match.group(1).strip() if ans_match else "",
                "solution": sol_match.group(1).strip() if sol_match else "",
                "key_points": kp_match.group(1).strip() if kp_match else ""
            })

    return parsed_questions


# ==========================================
# 4. सुंदर Word (.docx) फाइल तैयार करना
# ==========================================
def generate_clean_docx(mcqs, title="MCQ Question Bank"):
    doc = Document()
    
    # मुख्य शीर्षक
    title_p = doc.add_heading(title, level=1)
    title_p.alignment = 1

    doc.add_paragraph(f"कुल प्रश्न: {len(mcqs)}\n" + "—" * 45)

    for i, item in enumerate(mcqs, 1):
        # प्रश्न
        qp = doc.add_paragraph()
        q_run = qp.add_run(f"Q{i}. {item['question']}")
        q_run.bold = True
        q_run.font.size = Pt(11)

        # विकल्प
        for opt_label, opt_val in item['options']:
            op = doc.add_paragraph()
            op.paragraph_format.left_indent = Pt(18)
            op.add_run(f"({opt_label}) {opt_val.strip()}")

        # उत्तर
        if item['answer']:
            ap = doc.add_paragraph()
            ap.paragraph_format.left_indent = Pt(18)
            ans_run = ap.add_run(f"उत्तर: ({item['answer']})")
            ans_run.bold = True
            ans_run.font.color.rgb = RGBColor(16, 128, 64)

        # व्याख्या (Solution)
        if item['solution']:
            sp = doc.add_paragraph()
            sp.paragraph_format.left_indent = Pt(18)
            s_label = sp.add_run("व्याख्या: ")
            s_label.bold = True
            sp.add_run(item['solution'])

        # मुख्य बिंदु (Key Points)
        if item['key_points']:
            kp = doc.add_paragraph()
            kp.paragraph_format.left_indent = Pt(18)
            kp_label = kp.add_run("मुख्य बिंदु:\n")
            kp_label.bold = True
            kp.add_run(item['key_points'])

        doc.add_paragraph("—" * 35)

    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return out


# ==========================================
# 5. WeasyPrint से PDF तैयार करना
# ==========================================
def generate_clean_pdf(mcqs, title="MCQ Question Bank"):
    html_items = []
    
    for i, item in enumerate(mcqs, 1):
        opts_html = "".join([f'<div class="opt">({label}) {text.strip()}</div>' for label, text in item['options']])
        
        ans_html = f'<div class="ans"><b>उत्तर:</b> ({item["answer"]})</div>' if item['answer'] else ''
        sol_html = f'<div class="sol"><b>व्याख्या:</b> {item["solution"]}</div>' if item['solution'] else ''
        kp_html = f'<div class="kp"><b>मुख्य बिंदु:</b><br>{item["key_points"].replace(chr(10), "<br>")}</div>' if item['key_points'] else ''

        html_items.append(f"""
        <div class="card">
            <div class="q"><b>Q{i}.</b> {item['question']}</div>
            <div class="options">{opts_html}</div>
            {ans_html}
            {sol_html}
            {kp_html}
        </div>
        """)

    body_html = "\n".join(html_items)

    html_content = f"""
    <!DOCTYPE html>
    <html>
    <head>
        <meta charset="utf-8">
        <style>
            @page {{ size: A4; margin: 15mm; }}
            body {{ font-family: 'Helvetica Neue', Arial, sans-serif; font-size: 12px; color: #222; line-height: 1.5; }}
            h1 {{ text-align: center; color: #1e3a8a; font-size: 20px; border-bottom: 2px solid #cbd5e1; padding-bottom: 8px; }}
            .card {{ margin-bottom: 14px; padding-bottom: 10px; border-bottom: 1px dashed #94a3b8; page-break-inside: avoid; }}
            .q {{ font-weight: bold; font-size: 13px; color: #0f172a; margin-bottom: 6px; }}
            .options {{ margin-left: 15px; margin-bottom: 6px; }}
            .opt {{ margin-bottom: 2px; color: #334155; }}
            .ans {{ color: #15803d; font-weight: bold; margin-left: 15px; margin-top: 4px; }}
            .sol {{ margin-left: 15px; margin-top: 4px; color: #475569; }}
            .kp {{ margin-left: 15px; margin-top: 4px; background: #f8fafc; padding: 6px; border-left: 3px solid #3b82f6; font-size: 11px; }}
        </style>
    </head>
    <body>
        <h1>{title}</h1>
        <p style="text-align:center; color: #64748b;">कुल प्रश्न: {len(mcqs)}</p>
        {body_html}
    </body>
    </html>
    """

    out = io.BytesIO()
    HTML(string=html_content).write_pdf(out)
    out.seek(0)
    return out


# ==========================================
# 6. टेलीग्राम मैसेज हैंडलर्स
# ==========================================
@bot.message_handler(commands=['start', 'help'])
def send_welcome(message):
    msg = (
        "👋 *MCQ फ़ाइल कनवर्टर बॉट में आपका स्वागत है!*\n\n"
        "आप मुझे अपनी **.docx** (Word) फ़ाइल या **.txt** फ़ाइल भेज सकते हैं जिसमें इस तरह का फ़ॉर्मैट हो:\n"
        "- `Question:` ...\n"
        "- `(a)`, `(b)`, `(c)`, `(d)`\n"
        "- `Answer:` ...\n"
        "- `Solution:` ...\n\n"
        "मैं इसे साफ़-सुथरे **Word (.docx)** और **PDF** में तैयार करके वापस भेज दूंगा।"
    )
    bot.reply_to(message, msg, parse_mode="Markdown")


# फ़ाइल (DOCX / TXT) हैंडलर
@bot.message_handler(content_types=['document'])
def handle_incoming_file(message):
    file_name = message.document.file_name or "file.docx"
    file_ext = os.path.splitext(file_name)[1].lower()

    if file_ext not in ['.docx', '.txt']:
        bot.reply_to(message, "⚠️ कृपया केवल `.docx` (Word) या `.txt` फ़ाइल भेजें।")
        return

    status = bot.reply_to(message, f"⏳ `{file_name}` प्रोसेस की जा रही है...", parse_mode="Markdown")

    try:
        # फ़ाइल डाउनलोड करें
        file_info = bot.get_file(message.document.file_id)
        downloaded = bot.download_file(file_info.file_path)

        # टेक्स्ट एक्सट्रैक्ट करें
        if file_ext == '.docx':
            extracted_text = extract_text_from_docx(downloaded)
        else:
            extracted_text = downloaded.decode('utf-8', errors='ignore')

        # पार्सिंग
        mcqs = parse_mcqs(extracted_text)

        if not mcqs:
            bot.edit_message_text(
                "❌ फ़ाइल में कोई भी 'Question:' फ़ील्ड नहीं मिला। कृपया फ़ाइल का फ़ॉर्मैट चेक करें।",
                message.chat.id,
                status.message_id
            )
            return

        bot.edit_message_text(
            f"✅ {len(mcqs)} प्रश्न मिले! अब DOCX और PDF तैयार की जा रही है...",
            message.chat.id,
            status.message_id
        )

        base_title = os.path.splitext(file_name)[0]

        # 1. DOCX भेजें
        docx_data = generate_clean_docx(mcqs, title=base_title)
        docx_data.name = f"{base_title}_Formatted.docx"
        bot.send_document(message.chat.id, docx_data, caption=f"📄 {len(mcqs)} प्रश्नों की Word फ़ाइल")

        # 2. PDF भेजें
        pdf_data = generate_clean_pdf(mcqs, title=base_title)
        pdf_data.name = f"{base_title}_Formatted.pdf"
        bot.send_document(message.chat.id, pdf_data, caption=f"📕 {len(mcqs)} प्रश्नों की PDF फ़ाइल")

        bot.delete_message(message.chat.id, status.message_id)

    except Exception as e:
        bot.edit_message_text(f"❌ एरर: {str(e)}", message.chat.id, status.message_id)


# सीधे टेक्स्ट पेस्ट करने पर
@bot.message_handler(content_types=['text'])
def handle_incoming_text(message):
    if message.text.startswith('/'):
        return

    mcqs = parse_mcqs(message.text)
    if not mcqs:
        bot.reply_to(message, "⚠️ दिए गए टेक्स्ट में कोई 'Question:' नहीं मिला।")
        return

    status = bot.reply_to(message, f"⏳ {len(mcqs)} प्रश्न मिले, प्रोसेस किया जा रहा है...")

    try:
        docx_data = generate_clean_docx(mcqs, title="MCQ Bank")
        docx_data.name = "MCQs.docx"
        bot.send_document(message.chat.id, docx_data, caption="📄 आपकी Word फ़ाइल")

        pdf_data = generate_clean_pdf(mcqs, title="MCQ Bank")
        pdf_data.name = "MCQs.pdf"
        bot.send_document(message.chat.id, pdf_data, caption="📕 आपकी PDF फ़ाइल")

        bot.delete_message(message.chat.id, status.message_id)
    except Exception as e:
        bot.edit_message_text(f"❌ एरर: {str(e)}", message.chat.id, status.message_id)


# ==========================================
# 7. Flask और Webhook
# ==========================================
@app.route(f"/{BOT_TOKEN}", methods=['POST'])
def process_webhook():
    if request.headers.get('content-type') == 'application/json':
        json_data = request.get_data().decode('utf-8')
        update = telebot.types.Update.de_json(json_data)
        bot.process_new_updates([update])
        return 'OK', 200
    return 'Forbidden', 403


@app.route('/', methods=['GET', 'HEAD'])
@app.route('/ping', methods=['GET'])
def health():
    return jsonify({"status": "active", "service": "mcq-bot"}), 200


# सर्वर ऑन होते ही Webhook खुद सेट हो जाएगा
def setup_webhook():
    full_url = f"{WEBHOOK_URL.rstrip('/')}/{BOT_TOKEN}"
    try:
        info = bot.get_webhook_info()
        if info.url != full_url:
            bot.remove_webhook()
            bot.set_webhook(url=full_url)
            print(f"[+] Webhook registered: {full_url}")
    except Exception as ex:
        print(f"[-] Webhook setup failed: {ex}")

setup_webhook()

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 10000))
    app.run(host="0.0.0.0", port=port)
        
